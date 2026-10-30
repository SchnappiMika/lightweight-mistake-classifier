# -*- coding: utf-8 -*-
"""个性化分类结果导出模块（需求 4.3.2 第 5 点）。

两条导出通道：
    1. **Excel 分 sheet 归档**：按用户选定的分类维度（知识点 / 错误类型 / 难度 /
       复习优先级 / 任意自定义维度）把错题拆到不同 sheet，另附「全部错题」总表
       与「分类统计」汇总表，方便直接打印复习。
    2. **媒体分文件夹归档**：按分类标签把公式图片 / 图形题图片 / 听力音频复制到
       ``导出目录/<维度>/<标签名>/`` 下，并生成 index.csv 索引。

Windows 下 sheet 名与文件夹名有非法字符与长度限制，这里统一做安全化处理，
否则导出到一半会抛异常（对应测试项"导出的 Excel 格式规范、内容完整无错乱"）。
"""
from __future__ import annotations

import csv
import json
import logging
import math
import re
import shutil
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from app import config
from app.core.models import Question, Scheme

log = logging.getLogger(__name__)

# Excel sheet 名非法字符 + 长度上限 31
_RE_ILLEGAL_SHEET = re.compile(r"[\[\]\*\?:/\\]")
# Windows 文件夹名非法字符
_RE_ILLEGAL_DIR = re.compile(r'[<>:"/\\|\?\*\x00-\x1f]')
_WIN_RESERVED = {"CON", "PRN", "AUX", "NUL"} | {"COM{}".format(i) for i in range(1, 10)} \
    | {"LPT{}".format(i) for i in range(1, 10)}


def safe_sheet_name(name: str, used: Optional[set] = None) -> str:
    s = _RE_ILLEGAL_SHEET.sub("_", (name or "未分类").strip()) or "未分类"
    s = s[:31]
    if used is not None:
        base, i = s, 1
        while s in used:
            suffix = "_{}".format(i)
            s = base[:31 - len(suffix)] + suffix
            i += 1
        used.add(s)
    return s


def safe_dir_name(name: str) -> str:
    s = _RE_ILLEGAL_DIR.sub("_", (name or "未分类").strip()).rstrip(". ")
    if not s:
        s = "未分类"
    if s.upper() in _WIN_RESERVED:
        s = "_" + s
    return s[:80]


def default_export_path(prefix: str = "错题分类结果", ext: str = ".xlsx") -> Path:
    config.EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return config.EXPORT_DIR / "{}_{}{}".format(prefix, stamp, ext)


# --------------------------------------------------------------------------- #
# 行数据
# --------------------------------------------------------------------------- #
def scope_questions(questions: Sequence[Question], scheme: Optional[Scheme],
                    strict: bool = False) -> List[Question]:
    """Return questions belonging to the selected scheme.

    Older workspaces may contain records created before ``scheme_id`` was
    persisted.  If none of the supplied records has a scheme id, retain the
    legacy behaviour; once ids are present, an explicit scheme selection is
    always isolated.
    """
    items = list(questions)
    if scheme is None:
        return items
    has_ids = any(q.scheme_id for q in items)
    if strict or has_ids:
        return [q for q in items if q.scheme_id == scheme.id]
    return items


def _recognition_status(q: Question) -> str:
    statuses = []
    image_status = q.image_ocr_status
    if image_status:
        statuses.append("图片：{}".format(image_status))
    for attr, title in (("formula_ocr", "公式"), ("audio_transcript", "音频")):
        record = getattr(q, attr, None)
        if record is None:
            continue
        if getattr(record, "error", ""):
            state = "失败：{}".format(record.error)
        elif getattr(record, "confirmed", False):
            state = "已确认"
        elif getattr(record, "draft", ""):
            state = "待校对"
        else:
            state = "已识别"
        statuses.append("{}：{}".format(title, state))
    return "；".join(statuses)


def _row(q: Question, scheme: Optional[Scheme]) -> Dict[str, object]:
    row: Dict[str, object] = {
        "序号": "",
        "题干": q.raw_text or "",
        "学科": q.subject,
        "题型": config.QTYPE_LABELS.get(q.qtype, q.qtype),
        "人工标注": q.annotation or "",
        "媒体文件": Path(q.media_path).name if q.media_path else "",
        "媒体索引": q.id if q.media_path else "",
        "分类体系": scheme.name if scheme else "",
        "分类体系ID": q.scheme_id or (scheme.id if scheme else ""),
    }
    row["复习优先级标记"] = q.priority or ""
    row["识别状态"] = _recognition_status(q)
    # 保留旧列名，便于打开历史导出模板；新列汇总全部识别模态。
    row["图片识别状态"] = q.image_ocr_status
    row.update(_recognition_columns(q))
    row["错题ID"] = q.id
    row["耗时(ms)"] = round(q.elapsed_ms, 1)
    row["录入时间"] = datetime.fromtimestamp(q.created_at).strftime("%Y-%m-%d %H:%M")
    dims = scheme.dims if scheme else []
    reserved = set(row) | {"长文本明细"}
    for d in dims:
        if not d.enabled:
            continue
        title, suffix = d.name, 1
        while any(title + ending in reserved for ending in ("", "·置信度", "·人工修正")):
            title = "分类·{}{}".format(d.name, "" if suffix == 1 else "_{}".format(suffix))
            suffix += 1
        row[title] = q.label_names.get(d.key, "")
        row[title + "·置信度"] = round(float(q.scores.get(d.key, 0.0)), 3)
        row[title + "·人工修正"] = "是" if q.corrected.get(d.key) else ""
        reserved.update((title, title + "·置信度", title + "·人工修正"))
    return row


def _recognition_columns(q: Question) -> Dict[str, object]:
    """Keep machine output, editable drafts and actually applied text separate."""
    row = {}
    for attr, title in (("image_ocr", "图片识别"), ("formula_ocr", "公式识别"),
                        ("audio_transcript", "音频转写")):
        record = getattr(q, attr, None)
        for suffix, field in (("原始结果", "text"), ("校对草稿", "draft"),
                              ("已应用内容", "applied_text"), ("错误", "error"),
                              ("引擎", "engine"), ("耗时(ms)", "elapsed_ms"),
                              ("附件SHA256", "media_hash")):
            row[title + suffix] = getattr(record, field, "") if record else ""
        row[title + "已确认"] = ("是" if record.confirmed else "否") if record else ""
        row[title + "参数(JSON)"] = json.dumps(
            getattr(record, "recognition_options", {}), ensure_ascii=False) if record else ""
    row["图片识别当前题干"] = q.raw_text if q.image_ocr else ""
    row["图片识别行详情(JSON)"] = json.dumps(q.image_ocr.lines, ensure_ascii=False) if q.image_ocr else ""
    row["音频时长(s)"] = q.audio_transcript.duration_seconds if q.audio_transcript else ""
    feature = q.feature or {}
    row["图形提取参数(JSON)"] = json.dumps(feature.get("figure_parameters", []), ensure_ascii=False)
    row["预处理公式文本"] = feature.get("latex", "")
    row["预处理图片文字"] = feature.get("ocr_text", "")
    row["预处理听力文本"] = feature.get("transcript", "")
    media_titles = {"image": "图片文字", "formula": "公式", "audio": "听力音频"}
    row["分类所用媒体待校对"] = "、".join(
        media_titles.get(kind, kind) for kind in dict.fromkeys(feature.get("unreviewed_media", [])))
    row["本次媒体处理耗时(ms)"] = feature.get("media_elapsed_ms", "")
    row["本次自动处理总耗时(ms)"] = feature.get("total_process_ms", "")
    return row


def _display_width(text: str) -> int:
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


def _text_lines(text: str, width: float) -> int:
    return sum(max(1, math.ceil(_display_width(line) / max(1, width)))
               for line in text.split("\n"))


def _text_parts(text: str, width: int = 80, max_lines: int = 14):
    """Split without dropping whitespace; each part fits a printable detail row."""
    start = 0
    lines, column = 1, 0
    for index, char in enumerate(text):
        size = 0 if char in "\r\n" else _display_width(char)
        next_lines = lines + (char == "\n" or column + size > width)
        if next_lines > max_lines and index > start:
            yield text[start:index]
            start, lines, column = index, 1, 0
        if char == "\n":
            lines, column = lines + 1, 0
        elif column + size > width:
            lines, column = lines + 1, size
        else:
            column += size
    yield text[start:]


def _prepare_excel_rows(rows):
    details = []
    references = []
    for row in rows:
        fields = []
        for field, value in list(row.items()):
            if not isinstance(value, str) or (len(value) <= 32767 and _text_lines(value, 60) <= 16):
                continue
            fields.append(field)
            for part, text in enumerate(_text_parts(value), 1):
                details.append({"错题ID": row["错题ID"], "字段": field,
                                "分段": part, "全文分段": text})
            if len(value) > 32767:
                row[field] = value[:1000] + "\n（内容超过Excel单元格上限，全文见“长文本明细”）"
        references.append("、".join(fields))
    if details:
        for row, fields in zip(rows, references):
            row["长文本明细"] = "全文分段见“长文本明细”：" + fields if fields else ""
    return details


def _autofit(ws, df, max_width: int = 62) -> None:
    """按内容估算列宽（中文按 2 个字符宽计），让导出的表不用手动拉列。"""
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Alignment, Font, PatternFill

    for i, col in enumerate(df.columns, 1):
        width = _display_width(str(col))
        for v in df[col].astype(str).head(200):
            w = max((_display_width(line) for line in v.split("\n")), default=0)
            width = max(width, min(w, max_width - 2))
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), max_width)
    for cells in ws.iter_rows():
        height = 1
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, str):
                # Question text beginning with '=' is text, never an Excel formula.
                cell.data_type = "s"
                width = ws.column_dimensions[cell.column_letter].width - 2
                height = max(height, _text_lines(cell.value, width))
            if cell.row == 1:
                cell.font = Font(bold=True, color="17365D")
                cell.fill = PatternFill("solid", fgColor="E8F0FA")
        ws.row_dimensions[cells[0].row].height = min(300, max(30, height * 15 + 6))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.print_title_rows = "1:1"
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"


# --------------------------------------------------------------------------- #
# Excel 导出
# --------------------------------------------------------------------------- #
def export_excel(questions: Sequence[Question], scheme: Optional[Scheme],
                 path: Optional[str] = None,
                 split_dim: Optional[str] = config.DIM_KNOWLEDGE,
                 only_classified: bool = False,
                 strict_scheme: bool = False) -> str:
    """导出分类结果到 Excel。

    Args:
        split_dim: 用于分 sheet 的维度；``None`` 表示只出一张总表。
        only_classified: 只导出已完成分类的错题。
    Returns:
        实际写出的文件路径。
    """
    import pandas as pd

    scoped = scope_questions(questions, scheme, strict=strict_scheme)
    items = [q for q in scoped if (q.classified or not only_classified)]
    if not items:
        raise ValueError("没有可导出的错题")

    out = Path(path) if path else default_export_path()
    out.parent.mkdir(parents=True, exist_ok=True)

    rows = [_row(q, scheme) for q in items]
    for i, r in enumerate(rows, 1):
        r["序号"] = i
    detail_rows = _prepare_excel_rows(rows)
    df_all = pd.DataFrame(rows)

    dim_obj = scheme.dim(split_dim) if (scheme and split_dim) else None
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        # sheet 1：全部错题总表
        df_all.to_excel(writer, sheet_name="全部错题", index=False)
        _autofit(writer.sheets["全部错题"], df_all)
        if detail_rows:
            details = pd.DataFrame(detail_rows)
            details.to_excel(writer, sheet_name="长文本明细", index=False)
            _autofit(writer.sheets["长文本明细"], details, max_width=82)

        # sheet 2：分类统计汇总
        if scheme:
            stat_rows: List[Dict[str, object]] = []
            for d in scheme.dims:
                if not d.enabled:
                    continue
                counter: Dict[str, int] = {}
                for q in items:
                    counter[q.label_names.get(d.key, "未分类")] = \
                        counter.get(q.label_names.get(d.key, "未分类"), 0) + 1
                for name, cnt in sorted(counter.items(), key=lambda x: -x[1]):
                    stat_rows.append({"分类维度": d.name, "分类标签": name, "错题数": cnt,
                                      "占比": "{:.1%}".format(cnt / len(items))})
            if stat_rows:
                df_stat = pd.DataFrame(stat_rows)
                df_stat.to_excel(writer, sheet_name="分类统计", index=False)
                _autofit(writer.sheets["分类统计"], df_stat)

        # sheet 3..N：按维度分 sheet 归档
        if dim_obj is not None:
            used = {"全部错题", "分类统计", "长文本明细"}
            groups: Dict[str, List[Dict[str, object]]] = {}
            for q, r in zip(items, rows):
                groups.setdefault(q.label_names.get(split_dim, "") or "未分类", []).append(r)
            # 按体系中标签的定义顺序出 sheet，"未分类"排最后
            order = [lb.name for lb in dim_obj.labels] + ["未分类"]
            for name in sorted(groups, key=lambda n: order.index(n) if n in order else 999):
                sub = pd.DataFrame(groups[name])
                sn = safe_sheet_name("{}-{}".format(dim_obj.name, name), used)
                sub.to_excel(writer, sheet_name=sn, index=False)
                _autofit(writer.sheets[sn], sub)

    log.info("Excel 导出完成：%s（%d 条）", out, len(items))
    return str(out)


# --------------------------------------------------------------------------- #
# 媒体归档导出
# --------------------------------------------------------------------------- #
def export_media_archive(questions: Sequence[Question], scheme: Optional[Scheme],
                         folder: Optional[str] = None,
                         dim_key: str = config.DIM_KNOWLEDGE,
                         strict_scheme: bool = False) -> Tuple[str, int]:
    """按分类标签把图片 / 音频分文件夹归档，并生成 index.csv。

    Returns:
        （导出目录, 实际归档文件数）
    """
    scoped = scope_questions(questions, scheme, strict=strict_scheme)
    items = [q for q in scoped if q.media_path and Path(q.media_path).exists()]
    if not items:
        raise ValueError("没有可归档的图片或音频错题")

    dim_obj = scheme.dim(dim_key) if scheme else None
    dim_name = dim_obj.name if dim_obj else config.DIM_LABELS.get(dim_key, dim_key)

    root = Path(folder) if folder else default_export_path("错题媒体归档", "")
    root.mkdir(parents=True, exist_ok=True)

    index: List[Dict[str, object]] = []
    records = []
    n = 0
    for q in items:
        label = q.label_names.get(dim_key, "") or "未分类"
        sub = root / safe_dir_name(dim_name) / safe_dir_name(label)
        sub.mkdir(parents=True, exist_ok=True)
        src = Path(q.media_path)
        dst = sub / src.name
        i = 1
        while dst.exists():
            dst = sub / "{}_{}{}".format(src.stem, i, src.suffix)
            i += 1
        try:
            shutil.copy2(str(src), str(dst))
            n += 1
        except OSError as exc:
            log.warning("归档失败 %s: %s", src, exc)
            continue
        relative_path = str(dst.relative_to(root))
        entry = {
            "文件": relative_path,
            "分类标签": label,
            "学科": q.subject,
            "题型": config.QTYPE_LABELS.get(q.qtype, q.qtype),
            "题干": q.raw_text or "",
            "人工标注": q.annotation or "",
            "媒体索引": q.id,
            "分类体系": scheme.name if scheme else "",
            "分类体系ID": q.scheme_id or (scheme.id if scheme else ""),
            "识别状态": _recognition_status(q),
        }
        entry.update({key: value for key, value in _row(q, scheme).items() if key != "序号"})
        index.append(entry)
        records.append({"archive_file": relative_path, "question": q.to_dict()})

    if index:
        # utf-8-sig：让 Excel 双击打开 CSV 时中文不乱码
        with open(root / "index.csv", "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(index[0].keys()))
            w.writeheader()
            w.writerows(index)
        from app.core.storage import write_json
        write_json(root / "records.json", {
            "schema_version": 1,
            "scheme": scheme.to_dict() if scheme else None,
            "dimension": dim_key,
            "records": records,
        })
    log.info("媒体归档完成：%s（%d 个文件）", root, n)
    return str(root), n


# --------------------------------------------------------------------------- #
# 分类体系导出 / 导入（开源共享用）
# --------------------------------------------------------------------------- #
def export_scheme(scheme: Scheme, path: Optional[str] = None) -> str:
    """把一套分类体系导出为 JSON，便于同学之间共享词库。"""
    from app.core.storage import write_json
    out = Path(path) if path else default_export_path("分类体系_" + scheme.name, ".json")
    write_json(out, scheme.to_dict())
    return str(out)


def import_scheme(path: str) -> Scheme:
    """导入他人分享的分类体系 JSON。"""
    from app.core.models import Scheme as _S, new_id
    from app.core.storage import read_json
    data = read_json(Path(path), None)
    if not isinstance(data, dict):
        raise ValueError("不是合法的分类体系文件")
    sc = _S.from_dict(data)
    sc.id = new_id("sc")          # 换新 ID，避免与本地体系冲突
    sc.builtin = False
    return sc


def preview_structure(questions: Sequence[Question], scheme: Optional[Scheme],
                      split_dim: str, strict_scheme: bool = False) -> List[Tuple[str, int]]:
    """导出前的结构预览：返回 [(标签名, 错题数)]。"""
    counter: Dict[str, int] = {}
    for q in scope_questions(questions, scheme, strict=strict_scheme):
        counter[q.label_names.get(split_dim, "") or "未分类"] = \
            counter.get(q.label_names.get(split_dim, "") or "未分类", 0) + 1
    dim_obj = scheme.dim(split_dim) if scheme else None
    order = [lb.name for lb in dim_obj.labels] if dim_obj else []
    return sorted(counter.items(),
                  key=lambda x: order.index(x[0]) if x[0] in order else 999)
