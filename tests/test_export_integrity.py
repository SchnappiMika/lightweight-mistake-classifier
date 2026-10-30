from __future__ import annotations

import openpyxl

from app import config
from app.core.dataio import exporter
from app.core.models import CategoryLabel, Dimension, Question, Scheme


def _scheme(name: str, dim_key: str, dim_name: str) -> Scheme:
    return Scheme(
        name=name,
        subject=config.SUBJECT_MATH,
        dims=[Dimension(key=dim_key, name=dim_name,
                        labels=[CategoryLabel(id=name + "-1", name="标签A", dim=dim_key)])],
    )


def _question(scheme: Scheme, text: str) -> Question:
    question = Question(subject=config.SUBJECT_MATH, raw_text=text)
    question.scheme_id = scheme.id
    dim = scheme.dims[0]
    question.labels = {dim.key: dim.labels[0].id}
    question.label_names = {dim.key: dim.labels[0].name}
    question.classified = True
    return question


def test_export_isolated_by_scheme_and_keeps_actual_dimensions(tmp_path):
    first = _scheme("高数体系一", config.DIM_KNOWLEDGE, "知识点")
    second = _scheme("高数体系二", "mistake_source", "错因来源")
    q1 = _question(first, "第一套题")
    q2 = _question(second, "第二套题")

    output = exporter.export_excel([q1, q2], first, str(tmp_path / "first.xlsx"))
    workbook = openpyxl.load_workbook(output)
    rows = list(workbook["全部错题"].iter_rows(values_only=True))
    headers = list(rows[0])
    body = [dict(zip(headers, row)) for row in rows[1:]]

    assert [item["题干"] for item in body] == ["第一套题"]
    assert "知识点" in headers
    assert "错因来源" not in headers
    assert "分类体系" in headers and body[0]["分类体系"] == first.name
    assert "媒体索引" in headers and "识别状态" in headers

    output_second = exporter.export_excel([q1, q2], second, str(tmp_path / "second.xlsx"))
    rows_second = list(openpyxl.load_workbook(output_second)["全部错题"].iter_rows(values_only=True))
    headers_second = list(rows_second[0])
    body_second = [dict(zip(headers_second, row)) for row in rows_second[1:]]
    assert [row["题干"] for row in body_second] == ["第二套题"]
    assert "错因来源" in headers_second
    assert "知识点" not in headers_second


def test_scope_legacy_records_without_scheme_id_remains_compatible():
    question = Question(subject=config.SUBJECT_MATH, raw_text="旧记录")
    scheme = _scheme("旧体系", config.DIM_KNOWLEDGE, "知识点")
    assert exporter.scope_questions([question], scheme) == [question]
    assert exporter.scope_questions([question], scheme, strict=True) == []
