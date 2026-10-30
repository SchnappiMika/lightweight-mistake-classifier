"""Read actual exported files back to check complete media content and scope."""
from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl
import pytest

from app import config
from app.core.dataio import exporter
from app.core.models import (CategoryLabel, Dimension, FeatureBundle, ImageOcrRecord,
                             MediaReviewRecord, Question, Scheme)


def scheme(name="导出体系", dimension="知识点"):
    return Scheme(name=name, dims=[Dimension(key="topic", name=dimension,
                  labels=[CategoryLabel(name="积分复习", dim="topic")])])


def question(tmp_path, selected):
    attachment = tmp_path / (selected.id + ".png")
    attachment.write_bytes(b"original attachment bytes")
    item = Question(raw_text="题干，含逗号\n" + "完整题目" * 100,
                    annotation="人工标注\n" + "保留全部标注" * 90,
                    qtype=config.QTYPE_FORMULA, media_path=str(attachment),
                    scheme_id=selected.id, classified=True,
                    label_names={"topic": "积分复习"})
    item.image_ocr = ImageOcrRecord(text="原图文字", draft="图片校对草稿",
                                   applied_text="已应用图片文字", confirmed=True,
                                   engine="RapidOCR test", elapsed_ms=110,
                                   lines=[{"text": "原图文字", "score": 0.8, "box": [[1, 2]]}])
    item.formula_ocr = MediaReviewRecord(
        text=r"\int_0^1 x dx", draft=r"\int_0^1 x^2 dx",
        applied_text=r"\int_0^1 x dx", confirmed=True,
        engine="pix2tex", elapsed_ms=12345.67, media_hash="original-sha256",
        error="上次重试失败，保留已确认结果",
        recognition_options={"crop_box": [1, 2, 300, 120], "invert": True, "resize_mode": "fit"})
    item.audio_transcript = MediaReviewRecord(
        text="Original audio transcription.", draft="Edited audio draft.",
        applied_text="Previously confirmed audio.", confirmed=True,
        engine="whisper", elapsed_ms=2345.6, duration_seconds=18.25)
    item.feature = FeatureBundle(
        latex=r"\int_0^1 x dx", ocr_text="图中参数", transcript="Applied listening text.",
        figure_parameters=[{"name": "R", "value": "10", "unit": "kΩ",
                            "source": "annotation", "raw": "R = 10 kΩ"}]).to_dict()
    return item


def sheet_rows(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def test_excel_readback_keeps_original_draft_applied_metadata_and_parameter_source(tmp_path):
    selected = scheme()
    item = question(tmp_path, selected)
    path = exporter.export_excel([item], selected, str(tmp_path / "完整.xlsx"), split_dim="topic")
    workbook = openpyxl.load_workbook(path, data_only=False)
    for sheet_name in ("全部错题", "知识点-积分复习"):
        sheet = workbook[sheet_name]
        record = sheet_rows(sheet)[0]
        assert record["题干"] == item.raw_text and record["人工标注"] == item.annotation
        for attr, title in (("image_ocr", "图片识别"), ("formula_ocr", "公式识别"),
                            ("audio_transcript", "音频转写")):
            review = getattr(item, attr)
            assert record[title + "原始结果"] == review.text
            assert record[title + "校对草稿"] == review.draft
            assert record[title + "已应用内容"] == review.applied_text
            assert record[title + "引擎"] == review.engine
            assert record[title + "耗时(ms)"] == review.elapsed_ms
            assert record[title + "已确认"] == "是"
        assert record["公式识别错误"] == item.formula_ocr.error
        assert json.loads(record["公式识别参数(JSON)"]) == item.formula_ocr.recognition_options
        assert record["公式识别附件SHA256"] == "original-sha256"
        assert record["音频时长(s)"] == 18.25
        assert json.loads(record["图形提取参数(JSON)"]) == item.feature["figure_parameters"]
        assert json.loads(record["图片识别行详情(JSON)"]) == item.image_ocr.lines
        assert record["图片识别当前题干"] == item.raw_text
        assert record["预处理公式文本"] == item.feature["latex"]
        assert sheet.print_title_rows == "$1:$1" and sheet.freeze_panes == "A2"
        assert all(cell.alignment.wrap_text for cell in sheet[2])
        assert 30 <= sheet.row_dimensions[2].height <= 300


def test_confirmed_legacy_image_record_does_not_invent_previously_applied_text(tmp_path):
    selected = scheme()
    item = question(tmp_path, selected)
    item.image_ocr.applied_text = ""
    path = exporter.export_excel([item], selected, str(tmp_path / "legacy.xlsx"))
    row = sheet_rows(openpyxl.load_workbook(path)["全部错题"])[0]
    assert row["图片识别已应用内容"] is None
    assert row["图片识别当前题干"] == item.raw_text


def test_media_archive_csv_and_records_keep_full_contents_and_original_attachment(tmp_path):
    selected = scheme()
    item = question(tmp_path, selected)
    original = Path(item.media_path).read_bytes()
    directory, count = exporter.export_media_archive([item], selected,
                        str(tmp_path / "archive"), dim_key="topic", strict_scheme=True)
    root = Path(directory)
    assert count == 1 and (root / "index.csv").read_bytes().startswith(b"\xef\xbb\xbf")
    with (root / "index.csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    row = rows[0]
    assert row["题干"] == item.raw_text and row["人工标注"] == item.annotation
    assert row["公式识别原始结果"] == item.formula_ocr.text
    assert row["公式识别校对草稿"] == item.formula_ocr.draft
    assert row["公式识别已应用内容"] == item.formula_ocr.applied_text
    assert row["音频转写已应用内容"] == item.audio_transcript.applied_text
    assert json.loads(row["图形提取参数(JSON)"]) == item.feature["figure_parameters"]
    archive = json.loads((root / "records.json").read_text(encoding="utf-8"))
    assert archive["scheme"]["id"] == selected.id and archive["dimension"] == "topic"
    assert archive["records"][0]["question"] == item.to_dict()
    assert (root / archive["records"][0]["archive_file"]).read_bytes() == original
    assert Path(item.media_path).read_bytes() == original


def test_excel_long_text_details_reconstruct_beyond_cell_limit_without_loss(tmp_path):
    selected = scheme()
    item = question(tmp_path, selected)
    item.raw_text = "=题干不是Excel公式\n" + "很长的题目内容\n" * 5000
    item.formula_ocr.text = "\\frac{完整}{保留}" * 4000
    path = exporter.export_excel([item], selected, str(tmp_path / "超长.xlsx"))
    workbook = openpyxl.load_workbook(path, data_only=False)
    row = sheet_rows(workbook["全部错题"])[0]
    assert "长文本明细" in row["题干"]
    assert "题干" in row["长文本明细"] and "公式识别原始结果" in row["长文本明细"]
    details = sheet_rows(workbook["长文本明细"])
    for field, expected in (("题干", item.raw_text), ("公式识别原始结果", item.formula_ocr.text)):
        pieces = [entry for entry in details if entry["错题ID"] == item.id and entry["字段"] == field]
        assert "".join(entry["全文分段"] for entry in sorted(pieces, key=lambda value: value["分段"])) == expected
    assert workbook["长文本明细"].print_title_rows == "$1:$1"
    assert all(30 <= value.height <= 300 for value in workbook["长文本明细"].row_dimensions.values())


def test_equals_sign_text_remains_literal_after_excel_readback(tmp_path):
    selected = scheme()
    item = Question(raw_text="=SUM(A1:A2)", annotation="=1+2", scheme_id=selected.id)
    path = exporter.export_excel([item], selected, str(tmp_path / "literal.xlsx"))
    sheet = openpyxl.load_workbook(path, data_only=False)["全部错题"]
    row = sheet_rows(sheet)[0]
    assert row["题干"] == item.raw_text and row["人工标注"] == item.annotation
    assert sheet.cell(2, 2).data_type == "s"


def test_custom_dimension_names_do_not_overwrite_question_or_recognition_columns(tmp_path):
    selected = scheme()
    selected.dims = [Dimension(key="a", name="题干"),
                     Dimension(key="b", name="公式识别原始结果"),
                     Dimension(key="c", name="长文本明细")]
    item = question(tmp_path, selected)
    item.label_names = {"a": "维度甲", "b": "维度乙", "c": "维度丙"}
    path = exporter.export_excel([item], selected, str(tmp_path / "names.xlsx"))
    row = sheet_rows(openpyxl.load_workbook(path)["全部错题"])[0]
    assert row["题干"] == item.raw_text
    assert row["公式识别原始结果"] == item.formula_ocr.text
    assert row["分类·题干"] == "维度甲"
    assert row["分类·公式识别原始结果"] == "维度乙"
    assert row["分类·长文本明细"] == "维度丙"


@pytest.mark.parametrize("selected_index", [0, 1])
def test_both_excel_and_media_archive_isolate_two_schemes(tmp_path, selected_index):
    schemes = [scheme("体系甲", "知识点"), scheme("体系乙", "自定义掌握度")]
    questions = [question(tmp_path, entry) for entry in schemes]
    selected = schemes[selected_index]
    expected = questions[selected_index]
    path = exporter.export_excel(questions, selected, str(tmp_path / "isolated.xlsx"),
                                 strict_scheme=True)
    rows = sheet_rows(openpyxl.load_workbook(path)["全部错题"])
    assert len(rows) == 1 and rows[0]["错题ID"] == expected.id
    directory, count = exporter.export_media_archive(questions, selected,
                        str(tmp_path / "isolated"), dim_key="topic", strict_scheme=True)
    archive = json.loads((Path(directory) / "records.json").read_text(encoding="utf-8"))
    assert count == 1
    assert [entry["question"]["id"] for entry in archive["records"]] == [expected.id]
    assert archive["scheme"]["name"] == selected.name


@pytest.mark.parametrize("unreviewed, expected, media_ms", [
    (["image"], "图片文字", 1501.25),
    (["formula"], "公式", 25432.125),
    (["audio"], "听力音频", 12003.5),
    (["image", "formula", "audio"], "图片文字、公式、听力音频", 43210.5),
    (["formula"], "公式", 0),
    ([], None, 0),
])
def test_automatic_media_export_marks_unreviewed_and_keeps_distinct_timings(
        tmp_path, unreviewed, expected, media_ms):
    selected = scheme()
    item = question(tmp_path, selected)
    item.elapsed_ms = 8.76
    item.feature.update(unreviewed_media=unreviewed, media_elapsed_ms=media_ms,
                        total_process_ms=media_ms + 37.125)
    record_fields = {"image": ("image_ocr", "图片识别"),
                     "formula": ("formula_ocr", "公式识别"),
                     "audio": ("audio_transcript", "音频转写")}
    for kind in unreviewed:
        review = getattr(item, record_fields[kind][0])
        review.confirmed = False
        review.applied_text = ""
    output = exporter.export_excel([item], selected, str(tmp_path / "自动分类.xlsx"))
    row = sheet_rows(openpyxl.load_workbook(output)["全部错题"])[0]
    assert row["分类所用媒体待校对"] == expected
    assert row["本次媒体处理耗时(ms)"] == media_ms
    assert row["本次自动处理总耗时(ms)"] == media_ms + 37.125
    assert row["耗时(ms)"] == round(item.elapsed_ms, 1)
    for kind in unreviewed:
        title = record_fields[kind][1]
        assert row[title + "已确认"] == "否"
        assert row[title + "已应用内容"] is None
    directory, _ = exporter.export_media_archive([item], selected, str(tmp_path / "archive"))
    with (Path(directory) / "index.csv").open(encoding="utf-8-sig", newline="") as stream:
        csv_row = next(csv.DictReader(stream))
    assert csv_row["分类所用媒体待校对"] == (expected or "")
    assert float(csv_row["本次媒体处理耗时(ms)"]) == media_ms
    assert float(csv_row["本次自动处理总耗时(ms)"]) == media_ms + 37.125
    for kind in unreviewed:
        assert csv_row[record_fields[kind][1] + "已确认"] == "否"
    records = json.loads((Path(directory) / "records.json").read_text(encoding="utf-8"))
    assert records["records"][0]["question"]["feature"] == item.feature


def test_legacy_feature_does_not_claim_measured_automatic_media_timings(tmp_path):
    selected = scheme()
    item = question(tmp_path, selected)
    for key in ("unreviewed_media", "media_elapsed_ms", "total_process_ms"):
        item.feature.pop(key, None)
    path = exporter.export_excel([item], selected, str(tmp_path / "旧记录.xlsx"))
    row = sheet_rows(openpyxl.load_workbook(path)["全部错题"])[0]
    assert row["分类所用媒体待校对"] is None
    assert row["本次媒体处理耗时(ms)"] is None
    assert row["本次自动处理总耗时(ms)"] is None
