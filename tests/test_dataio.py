# -*- coding: utf-8 -*-
"""导入 / 导出模块测试（对应需求 4.4.1 第 4 点）。"""
from __future__ import annotations

from pathlib import Path

import pytest

from app import config
from app.core.dataio import exporter, importer
from app.core.models import CategoryLabel, Dimension, Question, Scheme


# --------------------------------------------------------------------------- #
# 导入
# --------------------------------------------------------------------------- #
class TestImportTxt:
    def test_blank_line_separated(self, tmp_path):
        p = tmp_path / "a.txt"
        p.write_text("求极限，用洛必达法则\n第二行还是同一条\n\n判断级数的敛散性\n",
                     encoding="utf-8")
        out = importer.import_txt(str(p))
        assert len(out) == 2
        assert "第二行还是同一条" in out[0].raw_text     # 段内换行不拆题

    def test_line_per_question_when_no_blank_line(self, tmp_path):
        p = tmp_path / "b.txt"
        p.write_text("求极限\n求积分\n求导数\n", encoding="utf-8")
        assert len(importer.import_txt(str(p))) == 3

    def test_comments_ignored(self, tmp_path):
        p = tmp_path / "c.txt"
        p.write_text("# 这是注释\n求极限\n\n# 又一条注释\n求积分\n", encoding="utf-8")
        out = importer.import_txt(str(p))
        assert len(out) == 2
        assert all("注释" not in q.raw_text for q in out)

    @pytest.mark.parametrize("enc", ["utf-8", "utf-8-sig", "gbk", "gb18030"])
    def test_chinese_encodings(self, tmp_path, enc):
        """Windows 记事本默认存 GBK，识别不了就等于用不了——这是高频真实场景。"""
        p = tmp_path / "enc.txt"
        p.write_bytes("求极限，用洛必达法则".encode(enc))
        out = importer.import_txt(str(p))
        assert len(out) == 1 and "洛必达法则" in out[0].raw_text

    def test_auto_infer_subject_and_qtype(self, tmp_path):
        p = tmp_path / "d.txt"
        p.write_text("用动量守恒定律求碰后速度\n\n" + r"求 $\int_0^1 x\,dx$" + "\n",
                     encoding="utf-8")
        out = importer.import_txt(str(p))
        assert out[0].subject == config.SUBJECT_PHYSICS
        assert out[1].qtype == config.QTYPE_FORMULA

    def test_missing_file_raises_friendly(self):
        with pytest.raises(importer.ImportError_, match="不存在"):
            importer.import_txt("/nope/none.txt")

    def test_empty_file_raises(self, tmp_path):
        p = tmp_path / "e.txt"
        p.write_text("\n\n  \n", encoding="utf-8")
        with pytest.raises(importer.ImportError_):
            importer.import_txt(str(p))

    def test_undecodable_file_raises(self, tmp_path):
        """所有候选编码都解不开时，要给一句人话，而不是抛原始的 UnicodeDecodeError。"""
        p = tmp_path / "bin.txt"
        p.write_bytes(bytes(range(1, 256)) * 3)
        with pytest.raises(importer.ImportError_, match="无法识别文件编码"):
            importer.import_txt(str(p))


class TestImportExcel:
    def _write(self, path, rows, columns=None):
        import pandas as pd
        pd.DataFrame(rows, columns=columns).to_excel(path, index=False)
        return str(path)

    def test_standard_columns(self, tmp_path):
        p = self._write(tmp_path / "a.xlsx", [
            {"题干": "求极限", "学科": "高等数学", "题型": "公式类", "标注": "洛必达用错"}])
        out = importer.import_excel(p)
        assert len(out) == 1
        q = out[0]
        assert q.subject == "高等数学" and q.qtype == config.QTYPE_FORMULA
        assert q.annotation == "洛必达用错"

    @pytest.mark.parametrize("col", ["题干", "题目", "错题", "内容", "question", "content"])
    def test_column_name_aliases(self, tmp_path, col):
        """用户的表头五花八门，认不出来就得让人手工改表——那就不叫"无需转换格式"了。"""
        p = self._write(tmp_path / "b.xlsx", [{col: "求极限"}])
        assert importer.import_excel(p)[0].raw_text == "求极限"

    @pytest.mark.parametrize("raw,expect", [
        ("公式类", config.QTYPE_FORMULA), ("formula", config.QTYPE_FORMULA),
        ("图形题", config.QTYPE_FIGURE), ("听力", config.QTYPE_LISTENING),
        ("阅读", config.QTYPE_READING), ("纯文本", config.QTYPE_TEXT),
    ])
    def test_qtype_aliases(self, tmp_path, raw, expect):
        p = self._write(tmp_path / "c.xlsx", [{"题干": "x", "题型": raw}])
        assert importer.import_excel(p)[0].qtype == expect

    def test_fallback_to_first_column(self, tmp_path):
        p = self._write(tmp_path / "d.xlsx", [{"随便什么表头": "求极限"}])
        assert importer.import_excel(p)[0].raw_text == "求极限"

    def test_skips_blank_rows(self, tmp_path):
        p = self._write(tmp_path / "e.xlsx",
                        [{"题干": "求极限"}, {"题干": ""}, {"题干": "求积分"}])
        assert len(importer.import_excel(p)) == 2

    def test_csv_supported(self, tmp_path):
        import pandas as pd
        p = tmp_path / "f.csv"
        pd.DataFrame([{"题干": "求极限"}]).to_csv(p, index=False, encoding="utf-8-sig")
        assert len(importer.import_excel(str(p))) == 1

    def test_sample_workbook_roundtrip(self):
        """项目自带的示例表必须能被自己导进来——否则"载入示例"按钮就是坏的。"""
        src = Path(__file__).resolve().parent.parent / "samples" / "sample_questions.xlsx"
        out = importer.import_excel(str(src))
        assert len(out) >= 150
        assert all(q.subject and q.qtype for q in out)


class TestImportSingle:
    def test_infers_qtype_from_latex(self):
        q = importer.import_single(r"求 $\lim_{x\to0}\frac{\sin x}{x}$")
        assert q.qtype == config.QTYPE_FORMULA
        assert q.source == "manual"
        # 题干里除了公式只剩一个"求"字，没有任何学科线索——此时判为"通用"是正确的，
        # 硬猜一个学科反而会把题分到错误的体系里去。
        assert q.subject == config.SUBJECT_GENERAL

    def test_infers_subject_when_text_has_clues(self):
        q = importer.import_single(r"用洛必达法则求 $\lim_{x\to0}\frac{\sin x}{x}$")
        assert q.subject == config.SUBJECT_MATH and q.qtype == config.QTYPE_FORMULA

    def test_explicit_beats_inference(self):
        q = importer.import_single("求极限", subject="大学物理", qtype=config.QTYPE_FIGURE)
        assert q.subject == "大学物理" and q.qtype == config.QTYPE_FIGURE


class TestScanDir:
    def test_splits_images_and_audio(self, tmp_path):
        (tmp_path / "a.png").write_bytes(b"x")
        (tmp_path / "b.mp3").write_bytes(b"x")
        (tmp_path / "c.txt").write_text("x")
        (tmp_path / "sub").mkdir()
        (tmp_path / "sub" / "d.jpg").write_bytes(b"x")
        imgs, auds = importer.scan_dir(str(tmp_path))
        assert len(imgs) == 2 and len(auds) == 1     # 递归扫描，且过滤掉 txt

    def test_missing_dir_is_empty(self):
        assert importer.scan_dir("/nope/none") == ([], [])


# --------------------------------------------------------------------------- #
# 导出
# --------------------------------------------------------------------------- #
@pytest.fixture
def sample_scheme():
    return Scheme(name="高数", subject=config.SUBJECT_MATH, dims=[
        Dimension(key=config.DIM_KNOWLEDGE, name="知识点", labels=[
            CategoryLabel(id="k1", name="定积分及其应用", dim=config.DIM_KNOWLEDGE),
            CategoryLabel(id="k2", name="无穷级数", dim=config.DIM_KNOWLEDGE)]),
        Dimension(key=config.DIM_ERROR, name="错误类型", labels=[
            CategoryLabel(id="e1", name="计算失误", dim=config.DIM_ERROR)]),
    ])


def make_q(text, kid, kname, ename="计算失误"):
    q = Question(subject=config.SUBJECT_MATH, qtype=config.QTYPE_TEXT, raw_text=text)
    q.labels = {config.DIM_KNOWLEDGE: kid, config.DIM_ERROR: "e1"}
    q.label_names = {config.DIM_KNOWLEDGE: kname, config.DIM_ERROR: ename}
    q.scores = {config.DIM_KNOWLEDGE: 0.83, config.DIM_ERROR: 0.5}
    q.classified = True
    return q


class TestSafeNames:
    @pytest.mark.parametrize("bad", ["a/b", "a\\b", "a*b", "a?b", "a[b]", "a:b"])
    def test_sheet_illegal_chars_removed(self, bad):
        """Excel 的 sheet 名不允许这些字符，不处理会导出到一半直接抛异常。"""
        out = exporter.safe_sheet_name(bad)
        assert not set(out) & set('[]*?:/\\')

    def test_sheet_length_capped(self):
        assert len(exporter.safe_sheet_name("知" * 60)) <= 31

    def test_sheet_dedup(self):
        used = set()
        a = exporter.safe_sheet_name("知识点-定积分", used)
        b = exporter.safe_sheet_name("知识点-定积分", used)
        assert a != b

    @pytest.mark.parametrize("bad", ['a<b', 'a>b', 'a"b', "a|b", "a/b"])
    def test_dir_illegal_chars_removed(self, bad):
        assert not set(exporter.safe_dir_name(bad)) & set('<>:"/\\|?*')

    @pytest.mark.parametrize("name", ["CON", "PRN", "NUL", "COM1", "LPT1"])
    def test_windows_reserved_names(self, name):
        """Windows 上不能有叫 CON/NUL 的文件夹，建目录会直接失败。"""
        assert exporter.safe_dir_name(name).upper() != name

    def test_trailing_dot_stripped(self):
        assert not exporter.safe_dir_name("知识点.").endswith(".")

    def test_empty_becomes_placeholder(self):
        assert exporter.safe_sheet_name("") == "未分类"
        assert exporter.safe_dir_name("   ") == "未分类"


class TestExportExcel:
    def test_sheets_structure(self, tmp_path, sample_scheme):
        import openpyxl
        qs = [make_q("求定积分", "k1", "定积分及其应用"),
              make_q("旋转体体积", "k1", "定积分及其应用"),
              make_q("判断敛散性", "k2", "无穷级数")]
        p = exporter.export_excel(qs, sample_scheme, str(tmp_path / "o.xlsx"))
        wb = openpyxl.load_workbook(p)
        assert "全部错题" in wb.sheetnames and "分类统计" in wb.sheetnames
        assert "知识点-定积分及其应用" in wb.sheetnames
        assert "知识点-无穷级数" in wb.sheetnames
        assert wb["知识点-定积分及其应用"].max_row == 3          # 表头 + 2 条
        assert wb["全部错题"].max_row == 4                       # 表头 + 3 条

    def test_all_dims_as_columns(self, tmp_path, sample_scheme):
        import openpyxl
        p = exporter.export_excel([make_q("求定积分", "k1", "定积分及其应用")],
                                  sample_scheme, str(tmp_path / "o.xlsx"))
        head = [c.value for c in openpyxl.load_workbook(p)["全部错题"][1]]
        for col in ("题干", "学科", "题型", "知识点", "知识点·置信度", "错误类型"):
            assert col in head

    def test_stats_sheet_content(self, tmp_path, sample_scheme):
        import openpyxl
        qs = [make_q("a", "k1", "定积分及其应用"), make_q("b", "k1", "定积分及其应用"),
              make_q("c", "k2", "无穷级数")]
        p = exporter.export_excel(qs, sample_scheme, str(tmp_path / "o.xlsx"))
        rows = list(openpyxl.load_workbook(p)["分类统计"].iter_rows(values_only=True))
        body = {r[1]: r[2] for r in rows[1:] if r[0] == "知识点"}
        assert body["定积分及其应用"] == 2 and body["无穷级数"] == 1

    def test_unclassified_goes_to_its_own_sheet(self, tmp_path, sample_scheme):
        import openpyxl
        q = Question(subject=config.SUBJECT_MATH, raw_text="没分类的题")
        q.classified = True
        p = exporter.export_excel([q], sample_scheme, str(tmp_path / "o.xlsx"))
        assert "知识点-未分类" in openpyxl.load_workbook(p).sheetnames

    def test_no_split_dim(self, tmp_path, sample_scheme):
        import openpyxl
        p = exporter.export_excel([make_q("a", "k1", "定积分及其应用")], sample_scheme,
                                  str(tmp_path / "o.xlsx"), split_dim=None)
        assert openpyxl.load_workbook(p).sheetnames == ["全部错题", "分类统计"]

    def test_only_classified_filter(self, tmp_path, sample_scheme):
        import openpyxl
        done = make_q("已分类", "k1", "定积分及其应用")
        todo = Question(subject=config.SUBJECT_MATH, raw_text="未分类")
        p = exporter.export_excel([done, todo], sample_scheme, str(tmp_path / "o.xlsx"),
                                  only_classified=True)
        assert openpyxl.load_workbook(p)["全部错题"].max_row == 2

    def test_empty_raises(self, sample_scheme):
        with pytest.raises(ValueError):
            exporter.export_excel([], sample_scheme)

    def test_label_with_illegal_char_does_not_break(self, tmp_path):
        """标签名里带斜杠（如"阅读/完形"）是很自然的写法，不能因此导出失败。"""
        import openpyxl
        sc = Scheme(name="英语", subject=config.SUBJECT_ENGLISH, dims=[
            Dimension(key=config.DIM_KNOWLEDGE, name="知识点", labels=[
                CategoryLabel(id="k1", name="阅读/完形*填空", dim=config.DIM_KNOWLEDGE)])])
        p = exporter.export_excel([make_q("x", "k1", "阅读/完形*填空")], sc,
                                  str(tmp_path / "o.xlsx"))
        wb = openpyxl.load_workbook(p)
        assert len(wb.sheetnames) == 3
        assert all(not set(n) & set('[]*?:/\\') for n in wb.sheetnames)


class TestExportMedia:
    def test_archives_by_label(self, tmp_path, sample_scheme):
        src = tmp_path / "src"
        src.mkdir()
        qs = []
        for i, (kid, kname) in enumerate([("k1", "定积分及其应用"), ("k1", "定积分及其应用"),
                                          ("k2", "无穷级数")]):
            f = src / "img{}.png".format(i)
            f.write_bytes(b"fake")
            q = make_q("题 {}".format(i), kid, kname)
            q.media_path = str(f)
            qs.append(q)
        root, n = exporter.export_media_archive(qs, sample_scheme, str(tmp_path / "out"))
        assert n == 3
        assert (Path(root) / "知识点" / "定积分及其应用").is_dir()
        assert len(list((Path(root) / "知识点" / "定积分及其应用").glob("*.png"))) == 2
        assert (Path(root) / "index.csv").exists()

    def test_index_csv_is_excel_friendly(self, tmp_path, sample_scheme):
        """index.csv 必须带 BOM，否则 Excel 双击打开中文全是乱码。"""
        f = tmp_path / "a.png"
        f.write_bytes(b"x")
        q = make_q("题", "k1", "定积分及其应用")
        q.media_path = str(f)
        root, _ = exporter.export_media_archive([q], sample_scheme, str(tmp_path / "out"))
        assert (Path(root) / "index.csv").read_bytes().startswith(b"\xef\xbb\xbf")

    def test_duplicate_filenames_not_overwritten(self, tmp_path, sample_scheme):
        qs = []
        for i in range(2):
            d = tmp_path / "d{}".format(i)
            d.mkdir()
            f = d / "same.png"                  # 同名不同目录
            f.write_bytes(b"x")
            q = make_q("题 {}".format(i), "k1", "定积分及其应用")
            q.media_path = str(f)
            qs.append(q)
        root, n = exporter.export_media_archive(qs, sample_scheme, str(tmp_path / "out"))
        assert n == 2
        assert len(list((Path(root) / "知识点" / "定积分及其应用").glob("*.png"))) == 2

    def test_no_media_raises(self, sample_scheme):
        with pytest.raises(ValueError):
            exporter.export_media_archive([make_q("无附件", "k1", "定积分及其应用")],
                                          sample_scheme)

    def test_missing_file_skipped(self, sample_scheme):
        q = make_q("题", "k1", "定积分及其应用")
        q.media_path = "/nope/gone.png"
        with pytest.raises(ValueError):          # 全部文件都不存在 -> 无可归档
            exporter.export_media_archive([q], sample_scheme)


class TestSchemeIO:
    def test_roundtrip_keeps_content(self, tmp_path, sample_scheme):
        p = exporter.export_scheme(sample_scheme, str(tmp_path / "s.json"))
        back = exporter.import_scheme(p)
        assert back.name == sample_scheme.name
        assert [d.key for d in back.dims] == [d.key for d in sample_scheme.dims]
        assert [lb.name for lb in back.dim(config.DIM_KNOWLEDGE).labels] == \
               [lb.name for lb in sample_scheme.dim(config.DIM_KNOWLEDGE).labels]

    def test_import_assigns_new_id(self, tmp_path, sample_scheme):
        """导入同学分享的体系时必须换新 ID，否则会和本地体系撞车。"""
        p = exporter.export_scheme(sample_scheme, str(tmp_path / "s.json"))
        back = exporter.import_scheme(p)
        assert back.id != sample_scheme.id and back.builtin is False

    def test_import_garbage_raises(self, tmp_path):
        p = tmp_path / "bad.json"
        p.write_text("[1,2,3]", encoding="utf-8")
        with pytest.raises(ValueError):
            exporter.import_scheme(str(p))


class TestPreview:
    def test_counts_and_order_follow_scheme(self, sample_scheme):
        qs = [make_q("a", "k2", "无穷级数"), make_q("b", "k1", "定积分及其应用"),
              make_q("c", "k1", "定积分及其应用")]
        out = exporter.preview_structure(qs, sample_scheme, config.DIM_KNOWLEDGE)
        assert out == [("定积分及其应用", 2), ("无穷级数", 1)]   # 按体系定义顺序，非出现顺序

    def test_unclassified_last(self, sample_scheme):
        q = Question(subject=config.SUBJECT_MATH, raw_text="x")
        out = exporter.preview_structure([q, make_q("a", "k1", "定积分及其应用")],
                                         sample_scheme, config.DIM_KNOWLEDGE)
        assert out[-1][0] == "未分类"
